from __future__ import annotations

import csv
import io
import json
import re
import unicodedata
from typing import Any

from openpyxl import load_workbook

from backend.app.pipeline.parsers.base import ParsedDocument, ParserError


def normalize_header(value: object) -> str:
    text = unicodedata.normalize("NFKD", str(value or ""))
    text = "".join(char for char in text if not unicodedata.combining(char))
    return re.sub(r"[^a-z0-9]+", " ", text.casefold()).strip()


def _parsed(records: list[dict[str, Any]], version: str) -> ParsedDocument:
    canonical = json.dumps(
        records, ensure_ascii=False, sort_keys=True, separators=(",", ":"), allow_nan=False
    )
    return ParsedDocument(
        structured_records=records,
        normalized_text="\n".join(
            " | ".join(f"{key}={value}" for key, value in record.items())
            for record in records
        ),
        canonical_json=canonical,
        parser_version=version,
    )


class IstatTerritoryParser:
    """Read an ISTAT XLSX without interpreting or persisting domain records."""

    version = "istat_territory_parser_v1"
    _required_header_groups = (
        {"codice regione", "cod regione"},
        {"denominazione in italiano", "denominazione comune", "comune"},
        {"codice comune formato alfanumerico", "codice comune", "cod comune"},
    )

    def parse(self, content: bytes) -> ParsedDocument:
        try:
            workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        except Exception as exc:
            raise ParserError(f"Invalid ISTAT XLSX: {exc}") from exc
        try:
            for sheet in workbook.worksheets:
                rows = sheet.iter_rows(values_only=True)
                for _ in range(25):
                    values = next(rows, None)
                    if values is None:
                        break
                    normalized = {normalize_header(value) for value in values if value is not None}
                    if all(group & normalized for group in self._required_header_groups):
                        headers = self._unique_headers(values)
                        records = [
                            {
                                header: self._cell(value)
                                for header, value in zip(headers, row)
                                if header and value is not None and str(value).strip()
                            }
                            for row in rows
                            if any(value is not None and str(value).strip() for value in row)
                        ]
                        if not records:
                            raise ParserError("ISTAT XLSX contains headers but no data rows")
                        return _parsed(records, self.version)
        finally:
            workbook.close()
        raise ParserError("ISTAT XLSX has no recognizable municipality header row")

    @staticmethod
    def _unique_headers(values: tuple[Any, ...]) -> list[str]:
        headers: list[str] = []
        seen: dict[str, int] = {}
        for index, value in enumerate(values):
            base = str(value).strip() if value is not None else ""
            if not base:
                headers.append("")
                continue
            seen[base] = seen.get(base, 0) + 1
            headers.append(base if seen[base] == 1 else f"{base} ({seen[base]})")
        return headers

    @staticmethod
    def _cell(value: Any) -> str | int | float | bool:
        if isinstance(value, float) and value.is_integer():
            return int(value)
        return value


class DaitMayorParser:
    """Parse DAIT's semicolon CSV into source-header-preserving row dictionaries."""

    version = "dait_mayor_parser_v1"

    def parse(self, content: bytes) -> ParsedDocument:
        try:
            text = content.decode("utf-8-sig")
        except UnicodeDecodeError:
            try:
                text = content.decode("cp1252")
            except UnicodeDecodeError as exc:
                raise ParserError(f"Invalid DAIT CSV encoding: {exc}") from exc
        body = self._table_body(text)
        reader = csv.DictReader(io.StringIO(body), delimiter=";")
        if not reader.fieldnames or len(reader.fieldnames) < 2:
            raise ParserError("DAIT CSV has no semicolon-delimited header")
        records = []
        for row in reader:
            record = {
                str(key).strip().strip('"'): value.strip()
                for key, value in row.items()
                if key is not None and value is not None and str(value).strip()
            }
            if record:
                records.append(record)
        if not records:
            raise ParserError("DAIT CSV contains no records")
        return _parsed(records, self.version)

    @staticmethod
    def _table_body(text: str) -> str:
        lines = text.splitlines()
        for index, line in enumerate(lines):
            cells = next(csv.reader([line], delimiter=";"))
            normalized = {normalize_header(cell) for cell in cells if str(cell).strip()}
            has_person = {"cognome", "nome"} <= normalized
            has_territory = any(
                "comune" in header or header == "sigla provincia"
                for header in normalized
            )
            if has_person and has_territory and len(normalized) >= 6:
                return "\n".join(lines[index:])
        raise ParserError("DAIT CSV has no recognizable mayor header row")
